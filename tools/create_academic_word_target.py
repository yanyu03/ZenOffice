"""Independent python-docx academic fixture; never uses the Zen writer."""
from pathlib import Path
from io import BytesIO
import json
import hashlib
import zipfile
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "RCD/red-team-corpus/academic-word-2026-10"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"


def field(p, instruction, cached="1"):
    r = p.add_run()._r
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); r.append(begin)
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = instruction; r.append(code)
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate"); r.append(separate)
    text = OxmlElement("w:t"); text.text = cached; r.append(text)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r.append(end)


def setup_section(sec, landscape=False, restart=None, roman=False):
    sec.different_first_page_header_footer = False
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
    sec.header.paragraphs[0].text = "Office 文档兼容性研究"
    sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, " PAGE ")
    for previous in list(sec._sectPr.findall(qn("w:pgNumType"))): sec._sectPr.remove(previous)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), "lowerRoman" if roman else "decimal")
    if restart is not None: pg.set(qn("w:start"), str(restart))
    sec._sectPr.append(pg)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.autofit = False
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    repeat = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, values): c.text = str(text)
    for row in t.rows:
        no_split = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(no_split)
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs: r.font.size = Pt(10.5)
    for cell in t.rows[0].cells:
        for r in cell.paragraphs[0].runs: r.bold = True
    return t


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "zen_academic_multiscenario.docx"
    if (OUT / "manifest.json").exists(): raise SystemExit("Frozen academic target exists; create a new corpus version instead of overwriting it.")
    doc = Document()
    for border in list(doc.styles.element.iter(qn("w:pBdr"))): border.getparent().remove(border)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption", "List Number"]:
        s = doc.styles[name]; s.font.name = "Times New Roman"; s.font.color.rgb = __import__('docx').shared.RGBColor(0,0,0)
        s.element.get_or_add_rPr().append(parse_xml(f'<w:rFonts xmlns:w="{W}" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/>'))
    doc.styles["Normal"].font.size = Pt(12)
    pf = doc.styles["Normal"].paragraph_format; pf.line_spacing = 1.35; pf.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(16); doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.core_properties.title = "Office 文档处理的内容与版面保真研究"
    doc.core_properties.author = "ZenOffice 工程验证"
    setup_section(doc.sections[0], restart=1, roman=True)
    doc.sections[0].different_first_page_header_footer = True
    doc.add_paragraph("Office 文档处理的内容与版面保真研究", "Title")
    doc.add_paragraph("学位论文格式与编辑保存验证样本", "Subtitle")
    doc.add_paragraph("ZenOffice 工程验证\n2026 年 10 月")
    doc.add_paragraph("本文件用于检验论文文档在读取、文字编辑和保存后的内容与版面保真。它覆盖常见论文结构，测试数据用于说明验收方法，不构成学术研究结论。")
    doc.add_paragraph().add_run("文档要求").bold = True
    doc.add_paragraph("编辑正文后，中文与英文字符、目录、页码、表格、图片、公式、脚注及各节方向应保持正确；无法安全执行的编辑必须明确拒绝。")
    doc.add_page_break()
    doc.add_paragraph("摘要", "Title")
    doc.add_paragraph("本文讨论 Office 文档引擎在内容修改与文件保存过程中需要保持的结构和版面条件。验证分为部件完整性、可见文本、字段与引用、原生宿主渲染四个层次。样本采用独立工具制作，避免用被测引擎自身生成的文件证明全部兼容能力。")
    doc.add_paragraph("关键词：文档保真；论文排版；受控编辑；验收")
    doc.add_paragraph().add_run("Abstract").bold = True
    doc.add_paragraph("This document defines an academic layout fixture for controlled editing and preservation checks. It combines sections, fields, tables, a figure, an equation and a footnote. The acceptance process compares actual content and native Office output before and after an edit.")
    doc.add_paragraph("Literal text: _x0001_ | café | naïve | 中文 | <tag> & ampersand.")
    doc.add_page_break()
    doc.add_paragraph("目录", "Title")
    field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', "目录将在 Word 中更新")
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE); setup_section(sec, restart=1)
    doc.add_paragraph("1 绪论", "Heading 1")
    doc.add_paragraph("1.1 问题与目标", "Heading 2")
    doc.add_paragraph("文档能够打开只是验收起点。可交付文件还应保持阅读顺序、分节方向、表格完整性和参考信息。若保存过程丢失了母版、页眉或字段，即使文字总数接近，也不能视为合格。")
    doc.add_paragraph("本研究样本将正文修改作为可观察操作，同时把其他部件的保持情况作为约束。这样既能证明文件实际发生了修改，也能检查未修改区域是否受到影响。")
    image = Image.new("RGB", (1200,240), "white"); draw = ImageDraw.Draw(image)
    for i, label in enumerate(["Reference", "Zen edit", "Native Office", "Compare"]):
        x = 12 + i*298; draw.rounded_rectangle((x,70,x+240,165), radius=12, outline="#333333", width=3)
        draw.text((x+26,108),label,fill="black",font_size=25)
        if i < 3: draw.line((x+243,118,x+284,118),fill="black",width=3)
    buf = BytesIO(); image.save(buf,format="PNG"); buf.seek(0)
    doc.add_picture(buf,width=Cm(15.5))
    p = doc.add_paragraph("图 ","Caption"); field(p," SEQ Figure "); p.add_run(" 文档编辑与验收流程")
    doc.add_paragraph("1.2 比较方法", "Heading 2")
    doc.add_paragraph("先冻结输入文件，再比较编辑前后的部件集合和可见内容。对于未修改的图片、字体、样式与关系，检查字节是否一致；对于发生修改的文本部件，检查修改位置之外的内容是否保持。")
    doc.add_page_break()
    doc.add_paragraph("2 方法", "Heading 1")
    doc.add_paragraph("2.1 内容指标", "Heading 2")
    doc.add_paragraph("定义保留比例为已保持内容项与应保持内容项的比值。该表达式只说明本样本采用的记账方法，不作为引擎通用正确性的证明。")
    equation = parse_xml('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>R = </m:t></m:r><m:f><m:num><m:r><m:t>N_preserved</m:t></m:r></m:num><m:den><m:r><m:t>N_required</m:t></m:r></m:den></m:f></m:oMath></m:oMathPara>')
    doc.add_paragraph()._p.append(equation)
    p = doc.add_paragraph("表 ","Caption"); field(p," SEQ Table "); p.add_run(" 受控样本的验收层次")
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"),"31"); start.set(qn("w:name"),"TableAcceptance")
    p._p.insert(0,start); end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"),"31"); p._p.append(end)
    table(doc,["层次","检查内容","通过条件"],[
        ["文件部件","图片、关系、样式与字段","未修改部件保持"],
        ["可见内容","正文、表格、引用与符号","内容正确且无遗漏"],
        ["原生排版","分页、页眉、表头与方向","页面完整可读"],
    ],[2.5,6,7.5])
    p = doc.add_paragraph("验证步骤见 "); field(p," REF TableAcceptance \\h ","表 1"); p.add_run("。")
    p = doc.add_paragraph("脚注与字段需要单独保持")
    r = p.add_run()._r; ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"),"1"); r.append(ref)
    p.add_run("，不能通过删除它们来让正文比较通过。")
    for text in ["冻结原始输入与来源信息。", "执行一次可观察的正文修改。", "由原生 Office 重开并检查内容与页面。"]:
        doc.add_paragraph(text,"List Number")
    doc.add_page_break()
    doc.add_paragraph("3 结果记录", "Heading 1")
    doc.add_paragraph("下表给出本样本需要保留的对象。记录中的状态描述验收条件，不宣称被测引擎已经通过。")
    table(doc,["对象","示例","验收要求"],[[i+1, name, "编辑后仍可读取并正确显示"] for i,name in enumerate([
        "中文正文","英文摘要","罗马数字页码","阿拉伯数字页码","自动目录","SEQ 题注","REF 引用","PNG 插图","OMML 公式","脚注","编号列表","横向附录"
    ])],[1.5,5.5,9])
    doc.add_paragraph("3.1 结果解释", "Heading 2")
    doc.add_paragraph("如果某项失败，应记录具体对象、页面和复现步骤，修复后使用相同输入再次比较。结果只适用于本次冻结的文件与操作，不扩大为任意论文的完全兼容结论。")
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE); setup_section(sec, landscape=True)
    doc.add_paragraph("附录 A 横向数据记录", "Heading 1")
    doc.add_paragraph("横向节用于验证页面方向、宽表格和页码连续性。下列数字是格式测试数据。")
    table(doc,["编号","样本名称","输入类型","编辑内容","保留对象","记录说明"],[[f"A{i:02d}","多场景论文","DOCX","正文词语替换","页眉、图片、字段","等待按冻结标准核验"] for i in range(1,7)],[1.4,4.0,2.5,4.5,6.0,6.0])
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE); setup_section(sec)
    doc.add_paragraph("参考资料", "Heading 1")
    doc.add_paragraph("[1] Southern Illinois University Graduate School. Templates. 官方论文模板页面，访问日期 2026-10-16。")
    p = doc.add_paragraph(); link = OxmlElement("w:hyperlink")
    rid = p.part.relate_to("https://siu.edu/admissions/graduate/academic-support/thesis-dissertation-research/etd/etd-templates.php", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    link.set(qn("r:id"),rid); run = OxmlElement("w:r"); txt = OxmlElement("w:t"); txt.text="SIU Graduate School Templates"; run.append(txt); link.append(run); p._p.append(link)
    doc.add_paragraph("[2] ZenOffice 工程验证。受控论文标的的对象清单与验收记录，2026。")
    doc.save(path)
    # Add a true footnote part independently of Zen's model.
    with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
    footnotes=f'<w:footnotes xmlns:w="{W}"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote><w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t> 脚注用于验证正文之外的关联部件在编辑保存后仍然存在。</w:t></w:r></w:p></w:footnote></w:footnotes>'
    parts['word/footnotes.xml']=footnotes.encode()
    from lxml import etree as E
    rels=E.fromstring(parts['word/_rels/document.xml.rels']); E.SubElement(rels,f'{{{R}}}Relationship',Id="rIdFootnotesZen",Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",Target="footnotes.xml")
    parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
    ct=E.fromstring(parts['[Content_Types].xml']); E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName="/word/footnotes.xml",ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
    parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items(): z.writestr(name,data)
    print(path)
    print(hashlib.sha256(path.read_bytes()).hexdigest())


if __name__ == "__main__": main()
